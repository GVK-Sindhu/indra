import os
import json
import csv
import subprocess
import sys
from datetime import datetime, timedelta

# Auto-install dependencies if missing
def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import fitz
except ImportError:
    install("PyMuPDF")
    import fitz

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    install("Pillow")
    from PIL import Image, ImageDraw, ImageFont

try:
    import openpyxl
except ImportError:
    install("openpyxl")
    import openpyxl

try:
    import docx
except ImportError:
    install("python-docx")
    import docx


BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "synthetic_indra_dataset"))
P101_DIR = os.path.join(BASE_DIR, "P-101")
GT_DIR = os.path.join(BASE_DIR, "ground_truth")

os.makedirs(P101_DIR, exist_ok=True)
os.makedirs(GT_DIR, exist_ok=True)

# Helper: Create text PDF
def create_text_pdf(filename, text, title=""):
    doc = fitz.open()
    page = doc.new_page()
    if title:
        page.insert_text(fitz.Point(50, 50), title, fontsize=16, fontname="helv", fontfile=None)
        page.insert_text(fitz.Point(50, 80), text, fontsize=11, fontname="helv", fontfile=None)
    else:
        page.insert_text(fitz.Point(50, 50), text, fontsize=11, fontname="helv", fontfile=None)
    doc.save(os.path.join(P101_DIR, filename))
    doc.close()

# Helper: Create scanned PDF (Image to PDF)
def create_scanned_pdf(filename, text):
    img = Image.new('RGB', (800, 1000), color=(240, 240, 230)) # yellowish paper
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 20)
    except:
        font = ImageFont.load_default()
    
    # Add some noise/scan artifacts lines
    for i in range(0, 1000, 20):
        d.line([(0, i), (800, i)], fill=(230, 230, 220), width=1)
        
    d.text((50, 50), text, fill=(50, 50, 50), font=font)
    
    pdf_path = os.path.join(P101_DIR, filename)
    img.save(pdf_path, "PDF", resolution=100.0)

# Helper: Create PNG image
def create_image(filename, text, size=(600, 300), bg_color=(200, 200, 200)):
    img = Image.new('RGB', size, color=bg_color)
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 24)
    except:
        font = ImageFont.load_default()
    
    d.text((20, 20), text, fill=(0, 0, 0), font=font)
    img.save(os.path.join(P101_DIR, filename))

# 1. P-101_OEM_Manual.pdf
create_text_pdf(
    "P-101_OEM_Manual.pdf",
    "ASSET TAG: P-101\n"
    "Model: FlowMaster X200\n"
    "Manufacturer: IndusPump Corp\n"
    "Type: Centrifugal Pump\n"
    "Rated Capacity: 500 m3/h\n"
    "Rated Pressure: 15 bar\n"
    "Max Allowable Working Pressure (MAWP): 16 bar\n"
    "OEM Recommendation: Standard maintenance interval every 8000 hours.\n"
    "Alert: High vibration exceeding 5.0 mm/s can cause catastrophic failure.",
    "OEM MANUAL - FlowMaster X200"
)

# 2. P-101_SOP_v1_Superseded.pdf (Contradiction 4.5 mm/s)
create_text_pdf(
    "P-101_SOP_v1_Superseded.pdf",
    "STANDARD OPERATING PROCEDURE\n"
    "Asset: Pump-101 (Alias for P-101)\n"
    "Version: 1.0 (SUPERSEDED)\n"
    "Date: 2021-05-15\n"
    "Status: OBSOLETE\n\n"
    "Operational Limits:\n"
    "- Maximum operating vibration threshold: 4.5 mm/s.\n"
    "- Shut down equipment if vibration exceeds this limit.",
    "P-101 SOP Version 1.0"
)

# 3. P-101_SOP_v2_Current.pdf (Contradiction 3.0 mm/s)
create_text_pdf(
    "P-101_SOP_v2_Current.pdf",
    "STANDARD OPERATING PROCEDURE\n"
    "Asset: P101 (Alias for P-101)\n"
    "Version: 2.0 (CURRENT)\n"
    "Date: 2025-01-10\n"
    "Status: APPROVED\n\n"
    "Operational Limits:\n"
    "- Maximum operating vibration threshold: 3.0 mm/s.\n"
    "- Shut down equipment immediately if vibration exceeds this limit to prevent bearing degradation.\n"
    "- Note: This supersedes all previous manuals due to updated safety guidelines.",
    "P-101 SOP Version 2.0"
)

# 4. P-101_Inspection_Warning.pdf (Native PDF - 3.8 mm/s)
create_text_pdf(
    "P-101_Inspection_Warning.pdf",
    "FIELD INSPECTION REPORT\n\n"
    "Date: 2026-06-12 10:00 AM\n"
    "Equipment Tag: P-101A\n"
    "Inspector: John Doe\n\n"
    "Findings:\n"
    "Visual inspection normal, but handheld sensor reads abnormal vibration.\n"
    "Measured Vibration: 3.8 mm/s.\n"
    "Warning: This exceeds the current SOP v2 limit of 3.0 mm/s.\n"
    "Action: Submitted work order WO-9921 for bearing investigation.",
    "P-101 FIELD INSPECTION REPORT"
)

# 5. P-101_Sensor_History.csv
csv_path = os.path.join(P101_DIR, "P-101_Sensor_History.csv")
with open(csv_path, 'w', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(["timestamp", "asset_id", "sensor", "parameter_value", "unit", "status"])
    writer.writerow(["2026-06-11T08:00:00Z", "P-101", "VIB-01", "1.2", "mm/s", "NORMAL"])
    writer.writerow(["2026-06-11T20:00:00Z", "P-101", "VIB-01", "2.1", "mm/s", "NORMAL"])
    writer.writerow(["2026-06-12T08:00:00Z", "P-101", "VIB-01", "3.2", "mm/s", "HIGH"])
    writer.writerow(["2026-06-12T09:30:00Z", "P-101", "VIB-01", "3.8", "mm/s", "CRITICAL"])

# 6. P-101_Maintenance_History.xlsx
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Maintenance History"
ws.append(["work_order_id", "asset_id", "date", "maintenance_type", "technician", "action_taken"])
ws.append(["WO-9921", "P-101", "2026-06-14", "Corrective", "Mike Smith", "Replaced degraded thrust bearing. Vibration returned to 1.1 mm/s."])
wb.save(os.path.join(P101_DIR, "P-101_Maintenance_History.xlsx"))

# 7. P-101_Incident_RCA.docx
doc = docx.Document()
doc.add_heading('Root Cause Analysis (RCA) - P-101 Bearing Failure', 0)
doc.add_paragraph('Date of Incident: 2026-06-12')
doc.add_paragraph('Asset: P-101')
doc.add_heading('Summary', level=1)
doc.add_paragraph('The centrifugal pump P-101 experienced severe bearing degradation requiring replacement via WO-9921.')
doc.add_heading('Root Cause', level=1)
doc.add_paragraph('The unit was operated at 3.8 mm/s vibration for over 24 hours. While this was below the old superseded v1 limit of 4.5 mm/s, it exceeded the current approved v2 limit of 3.0 mm/s. The failure to adhere to the current SOP v2 resulted in accelerated wear.')
doc.save(os.path.join(P101_DIR, "P-101_Incident_RCA.docx"))

# 8. P-101_Asset_Metadata.json
meta = {
    "asset_id": "P-101",
    "equipment_type": "Centrifugal Pump",
    "aliases": ["P101", "Pump-101", "P-101A"],
    "location": "Unit 3",
    "criticality": "High",
    "status": "Active"
}
with open(os.path.join(P101_DIR, "P-101_Asset_Metadata.json"), "w") as f:
    json.dump(meta, f, indent=4)

# 9. P-101_Nameplate.png
create_image(
    "P-101_Nameplate.png",
    "IndusPump Corp\nModel: FlowMaster X200\nS/N: 99887766\nTAG: P-101\nRated: 15 bar",
    size=(400, 200),
    bg_color=(192, 192, 192)
)

# 10. P-101_PID_Snippet.png
create_image(
    "P-101_PID_Snippet.png",
    "P&ID DIAGRAM SNIPPET\n\n[Tank T-100] ---> ( Valve V-01 ) ---> [ P-101 ] ---> ( Flow Meter FT-101 )\n\nTags: T-100, V-01, P-101, FT-101",
    size=(800, 200),
    bg_color=(255, 255, 255)
)

# ---------------- GROUND TRUTH ----------------

# manifest.json
manifest = [
    {"file_name": "P-101_OEM_Manual.pdf", "asset_id": "P-101", "document_type": "Manual", "file_format": "PDF"},
    {"file_name": "P-101_SOP_v1_Superseded.pdf", "asset_id": "P-101", "document_type": "SOP", "file_format": "PDF", "revision": "1.0", "approval_status": "OBSOLETE"},
    {"file_name": "P-101_SOP_v2_Current.pdf", "asset_id": "P-101", "document_type": "SOP", "file_format": "PDF", "revision": "2.0", "approval_status": "APPROVED"},
    {"file_name": "P-101_Inspection_Warning.pdf", "asset_id": "P-101", "document_type": "Inspection", "file_format": "PDF"},
    {"file_name": "P-101_Sensor_History.csv", "asset_id": "P-101", "document_type": "Telemetry", "file_format": "CSV"},
    {"file_name": "P-101_Maintenance_History.xlsx", "asset_id": "P-101", "document_type": "Maintenance", "file_format": "XLSX"},
    {"file_name": "P-101_Incident_RCA.docx", "asset_id": "P-101", "document_type": "RCA", "file_format": "DOCX"},
    {"file_name": "P-101_Asset_Metadata.json", "asset_id": "P-101", "document_type": "Metadata", "file_format": "JSON"},
    {"file_name": "P-101_Nameplate.png", "asset_id": "P-101", "document_type": "Image", "file_format": "PNG"},
    {"file_name": "P-101_PID_Snippet.png", "asset_id": "P-101", "document_type": "Image", "file_format": "PNG"}
]
with open(os.path.join(GT_DIR, "dataset_manifest.json"), "w") as f:
    json.dump(manifest, f, indent=4)

# contradictions.json
contradictions = [{
    "contradiction_id": "C-001",
    "asset_id": "P-101",
    "fact_or_parameter": "maximum operating vibration threshold",
    "source_document_1": "P-101_SOP_v1_Superseded.pdf",
    "value_1": "4.5 mm/s",
    "source_document_2": "P-101_SOP_v2_Current.pdf",
    "value_2": "3.0 mm/s",
    "reason_for_conflict": "SOP version update",
    "expected_preferred_source": "P-101_SOP_v2_Current.pdf",
    "explanation": "Version 2.0 is marked APPROVED and explicitly supersedes the OBSOLETE version 1.0."
}]
with open(os.path.join(GT_DIR, "contradictions.json"), "w") as f:
    json.dump(contradictions, f, indent=4)

# evidence_chains.json
chains = [{
    "asset_id": "P-101",
    "chain_name": "Vibration Incident to Bearing Replacement",
    "steps": [
        {"step": 1, "document": "P-101_Sensor_History.csv", "finding": "Vibration rises to 3.8 mm/s"},
        {"step": 2, "document": "P-101_Inspection_Warning.pdf", "finding": "Inspector notes 3.8 mm/s exceeds v2 limit, opens WO-9921"},
        {"step": 3, "document": "P-101_Maintenance_History.xlsx", "finding": "WO-9921 completed: bearing replaced"},
        {"step": 4, "document": "P-101_Incident_RCA.docx", "finding": "RCA determines failure caused by operating above 3.0 mm/s limit"}
    ]
}]
with open(os.path.join(GT_DIR, "evidence_chains.json"), "w") as f:
    json.dump(chains, f, indent=4)

# benchmark_questions.json
qs = [
    {"question_id": 1, "asset_id": "P-101", "question": "What is the rated pressure for P-101?", "reasoning_type": "single-document", "expected_answer": "15 bar", "required_sources": ["P-101_OEM_Manual.pdf"]},
    {"question_id": 2, "asset_id": "P-101", "question": "What is the manufacturer of P-101?", "reasoning_type": "cross-format", "expected_answer": "IndusPump Corp", "required_sources": ["P-101_Nameplate.png", "P-101_OEM_Manual.pdf"]},
    {"question_id": 3, "asset_id": "P-101", "question": "What was the vibration reading on 2026-06-12 at 09:30?", "reasoning_type": "single-document", "expected_answer": "3.8 mm/s", "required_sources": ["P-101_Sensor_History.csv"]},
    {"question_id": 4, "asset_id": "P-101", "question": "Who was the inspector that submitted WO-9921?", "reasoning_type": "single-document", "expected_answer": "John Doe", "required_sources": ["P-101_Inspection_Warning.pdf"]},
    {"question_id": 5, "asset_id": "P-101", "question": "What action was taken in WO-9921?", "reasoning_type": "maintenance-history", "expected_answer": "Replaced degraded thrust bearing", "required_sources": ["P-101_Maintenance_History.xlsx"]},
    {"question_id": 6, "asset_id": "P-101", "question": "What is the current maximum operating vibration threshold for P-101?", "reasoning_type": "contradiction-detection", "expected_answer": "3.0 mm/s", "required_sources": ["P-101_SOP_v2_Current.pdf"]},
    {"question_id": 7, "asset_id": "P-101", "question": "Why did the bearing fail according to the RCA?", "reasoning_type": "cross-document", "expected_answer": "Operated above the current 3.0 mm/s limit", "required_sources": ["P-101_Incident_RCA.docx", "P-101_SOP_v2_Current.pdf"]},
    {"question_id": 8, "asset_id": "P-101", "question": "Are P101 and Pump-101 the same asset?", "reasoning_type": "asset-alias", "expected_answer": "Yes", "required_sources": ["P-101_Asset_Metadata.json"]},
    {"question_id": 9, "asset_id": "P-101", "question": "What is the safe operating temperature for P-101?", "reasoning_type": "unsupported-abstention", "expected_answer": "Abstain (No data available)", "required_sources": []},
    {"question_id": 10, "asset_id": "P-101", "question": "How many hours of operation until the next overhaul?", "reasoning_type": "unsupported-abstention", "expected_answer": "Abstain (No data available)", "required_sources": []}
]
with open(os.path.join(GT_DIR, "benchmark_questions.json"), "w") as f:
    json.dump(qs, f, indent=4)

print(f"Generated 10 P-101 files and 4 ground truth files in {BASE_DIR}")
